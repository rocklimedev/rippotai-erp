"""Read the supplied reference; generate versioned planner defaults, never sample project data."""
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
w = openpyxl.load_workbook(root / 'documentation/templates/v1/PROEJCT PLANNER VF.xlsx')
tasks = []
for sheet, module in [('Consultancy', 'CONSULTANCY'), ('PMC', 'PMC')]:
    phase = None
    merged_work = None
    for row in range(5, 55):
        s = w[sheet]
        phase = s.cell(row, 2).value or phase
        work = s.cell(row, 3).value
        details = s.cell(row, 4).value
        if not work and not details:
            continue
        # Only propagate work labels that are explicitly merged in the source.
        for area in s.merged_cells.ranges:
            if area.min_col == 3 and area.min_row <= row <= area.max_row:
                work = s.cell(area.min_row, 3).value
        tasks.append(dict(module=module, phase=phase.strip(), work_name=(work or '').strip(), details=(details or '').strip(), sort_order=row-5))
procurement = []
for row in range(5, 27):
    labour, material = w['Vendor & Procurement'].cell(row, 2).value, w['Vendor & Procurement'].cell(row, 3).value
    if labour or material:
        procurement.append(dict(item_type='LABOUR' if labour else 'MATERIAL', category_name=(labour or material).strip(), sort_order=row-5))
(root / 'backend/src/modules/projects/planner-workbook-template.ts').write_text('export const WORKBOOK_TEMPLATE = ' + json.dumps(dict(tasks=tasks, procurement=procurement), ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
print(f'{len(tasks)} tasks; {len(procurement)} procurement categories')
